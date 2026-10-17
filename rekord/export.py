"""Exports: Markdown, plain text, SRT, WebVTT, JSON, DOCX (if python-docx installed)."""
from __future__ import annotations

import datetime as dt
import io
import json

from . import db
from .summarize import fmt_time


def _ts(t: float, sep: str) -> str:
    ms = int(round(t * 1000))
    return f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:{ms % 60000 // 1000:02d}{sep}{ms % 1000:03d}"


def _data(rid: str):
    rec = db.get_recording(rid)
    sp = {s["id"]: s["name"] for s in db.list_speakers(rid)}
    segs = db.list_segments(rid)
    return rec, sp, segs


def markdown(rid: str, with_summary: bool = True) -> str:
    rec, sp, segs = _data(rid)
    date = dt.datetime.fromtimestamp(rec["created_at"]).strftime("%d.%m.%Y %H:%M")
    out = [f"# {rec['title']}", "", f"*{date} · {fmt_time(rec['duration'] or 0)} · {', '.join(dict.fromkeys(sp.values()))}*", ""]
    if with_summary:
        for s in db.list_summaries(rid):
            out += [f"## {s['title']}", "", s["content"], ""]
    if rec.get("notes"):
        out += ["## Notizen", "", rec["notes"], ""]
    out += ["## Transkript", ""]
    for s in segs:
        out.append(f"**{sp.get(s['speaker_id'], '?')}** [{fmt_time(s['start'])}]: {s['text']}  ")
    return "\n".join(out) + "\n"


def text(rid: str) -> str:
    _, sp, segs = _data(rid)
    return "\n".join(f"[{fmt_time(s['start'])}] {sp.get(s['speaker_id'], '?')}: {s['text']}" for s in segs) + "\n"


def srt(rid: str) -> str:
    _, sp, segs = _data(rid)
    return "\n".join(f"{i}\n{_ts(s['start'], ',')} --> {_ts(s['end'], ',')}\n{sp.get(s['speaker_id'], '?')}: {s['text']}\n"
                     for i, s in enumerate(segs, 1))


def vtt(rid: str) -> str:
    _, sp, segs = _data(rid)
    return "WEBVTT\n\n" + "\n".join(f"{_ts(s['start'], '.')} --> {_ts(s['end'], '.')}\n<v {sp.get(s['speaker_id'], '?')}>{s['text']}\n"
                                    for s in segs)


def as_json(rid: str) -> str:
    rec, sp, segs = _data(rid)
    return json.dumps({"recording": rec | {"peaks": None}, "speakers": db.list_speakers(rid),
                       "segments": [s | {"speaker": sp.get(s["speaker_id"])} for s in segs],
                       "summaries": db.list_summaries(rid)}, ensure_ascii=False, indent=2)


def docx(rid: str) -> bytes:
    from docx import Document  # optional dependency
    rec, sp, segs = _data(rid)
    d = Document()
    d.add_heading(rec["title"], 0)
    for s in db.list_summaries(rid):
        d.add_heading(s["title"], 1)
        for line in s["content"].splitlines():
            if line.startswith("#"):
                d.add_heading(line.lstrip("# "), 2)
            elif line.strip():
                d.add_paragraph(line.strip("- "), style="List Bullet" if line.lstrip().startswith("-") else None)
    d.add_heading("Transkript", 1)
    for s in segs:
        p = d.add_paragraph()
        p.add_run(f"{sp.get(s['speaker_id'], '?')} [{fmt_time(s['start'])}]: ").bold = True
        p.add_run(s["text"])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


FORMATS = {
    "md": ("text/markdown", markdown),
    "txt": ("text/plain", text),
    "srt": ("application/x-subrip", srt),
    "vtt": ("text/vtt", vtt),
    "json": ("application/json", as_json),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx),
}
